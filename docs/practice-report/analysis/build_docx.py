"""Збирає section-NN.md у .docx за форматом шаблону звіту (TNR 14, A4, поля 2,54 см)."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")


def set_font(run, size=14, bold=None, italic=None, mono=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if mono:
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_inline(par, text, size=14, bold=False, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); set_font(r, size, True, italic)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); set_font(r, size - 2, bold, italic, mono=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, size, bold, True)
        else:
            r = par.add_run(part); set_font(r, size, bold, italic)


def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, before=0, after=0, keep=False):
    pf = par.paragraph_format
    pf.alignment = align
    pf.line_spacing = 1.5
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Cm(1.25) if indent else None
    pf.keep_with_next = keep


def shade(cell, hexcolor):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def borders(table, dashed=False):
    tblPr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "dashed" if dashed else "single")
        e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "808080")
        b.append(e)
    tblPr.append(b)


def page_number_footer(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); set_font(r, 12)
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), t)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        r._r.append(el)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, row in enumerate([header] + body):
        for j, txt in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            add_inline(p, txt.strip(), size=11, bold=(i == 0))
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            if i == 0:
                shade(c, "E7E6E6")
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(6)


def add_figure(doc, block):
    text = "\n".join(l[1:].strip() for l in block)
    cap = re.search(r"\*Підпис:\*\s*(.+)", text)
    desc = re.search(r"\*Що зобразити:\*\s*(.+)", text, re.S)
    num = re.search(r"РИСУНОК ([\d.]+)", text)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, dashed=True)
    c = t.cell(0, 0); shade(c, "F5F5F5")
    c.text = ""
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"\n\n[ Місце для рисунка {num.group(1) if num else ''} ]\n"); set_font(r, 13, True, color=(0x80, 0x80, 0x80))
    if desc:
        p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run("Вказівка: " + desc.group(1).strip() + "\n\n"); set_font(r2, 10, italic=True, color=(0x80, 0x80, 0x80))
    cp = doc.add_paragraph()
    add_inline(cp, cap.group(1).strip() if cap else "", size=14)
    fmt(cp, WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=6, after=12)


def build(src: Path, out: Path):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    page_number_footer(doc)

    lines = src.read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip() or ln.strip() == "---":
            i += 1; continue
        if ln.startswith("# "):
            p = doc.add_paragraph(); add_inline(p, ln[2:], bold=True); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, False, 0, 12, True)
        elif ln.startswith("## "):
            p = doc.add_paragraph(); add_inline(p, ln[3:], bold=True); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, False, 12, 6, True)
        elif ln.startswith("### "):
            p = doc.add_paragraph(); add_inline(p, ln[4:], bold=True, italic=False); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, False, 6, 3, True)
        elif ln.startswith("$$ "):
            p = doc.add_paragraph(); add_inline(p, ln[3:].strip())
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 6)
        elif ln.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.+)\)", ln)
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 0, True)
            p.paragraph_format.line_spacing = 1.0
            p.add_run().add_picture(str(src.parent / m.group(2)), width=Cm(15.5))
            cp = doc.add_paragraph(); add_inline(cp, m.group(1))
            fmt(cp, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 12)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"\s*-+\s*", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows); continue
        elif ln.startswith(">"):
            block = []
            while i < len(lines) and lines[i].startswith(">"):
                block.append(lines[i]); i += 1
            add_figure(doc, block); continue
        elif re.match(r"^(\d+)\.\s", ln) or ln.startswith("- "):
            m = re.match(r"^(\d+\.)\s(.*)", ln)
            p = doc.add_paragraph()
            if m:
                add_inline(p, f"{m.group(1)} {m.group(2)}")
            else:
                add_inline(p, "– " + ln[2:])
            fmt(p, indent=False, after=0)
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-0.6)
        else:
            p = doc.add_paragraph(); add_inline(p, ln)
            is_caption = ln.startswith("**Таблиця")
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT if is_caption else WD_ALIGN_PARAGRAPH.JUSTIFY, not is_caption, 6 if is_caption else 0, 3 if is_caption else 0, is_caption)
        i += 1
    doc.save(out)


if __name__ == "__main__":
    src = Path(sys.argv[1]); build(src, src.with_suffix(".docx")); print(src.with_suffix(".docx"))
