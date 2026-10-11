"""Збирає звіт із Markdown-розділів у .md і .docx. Дві версії:

    python docs/practice-report/analysis/build_report.py practice   # report/ -> practice-report.docx (~90 с.)
    python docs/practice-report/analysis/build_report.py full       # thesis-full/ -> thesis-full.docx


Формат: A4, Times New Roman 14, інтервал 1,5, береги 25/15/20/20 мм (ліве/праве/верхнє/нижнє).

Запуск (потрібні python-docx і Pillow):
    python docs/practice-report/analysis/build_report.py

Посилання в тексті записано ключами [@key]; скрипт нумерує джерела в порядку першої
появи, додає перелік джерел і перевіряє, що на умовну сторінку тексту (PAGE_CHARS
символів) припадає не більше MAX_SOURCES_PER_PAGE різних джерел.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(Path(__file__).parent))
import build_docx as bd  # noqa: E402

HERE = Path(__file__).parent
VARIANTS = {"practice": ("report", "practice-report"), "full": ("thesis-full", "thesis-full")}
# У звіті з практики розділи йдуть у порядку завдань: інтеграція з Meet (файл 06) перед взаємодією з LLM (файл 05).
PRACTICE_PARTS = ["00-intro.md", "02-analysis.md", "03-design.md", "04-algorithms.md", "06-extension.md",
                  "05-llm.md", "07-evaluation.md", "08-conclusions.md"]
REPORT = HERE.parent / "report"
PARTS = ["00-intro.md", "02-analysis.md", "03-design.md", "04-algorithms.md", "05-llm.md",
         "06-extension.md", "07-evaluation.md", "08-conclusions.md"]
ACCESSED_OLD, ACCESSED_PRICES, ACCESSED_NEW = "05.10.2026", "08.10.2026", "09.10.2026"
PAGE_CHARS, MAX_SOURCES_PER_PAGE = 1800, 5
MAX_IMAGE_WIDTH_CM, MAX_IMAGE_HEIGHT_CM = 13.5, 10.5
TABLE_FONT_PT = 10


def url(u: str, date: str) -> str:
    return f"URL: {u} (дата звернення: {date})."


REFS: dict[str, str] = {
    "wooldridge": "Wooldridge M. An Introduction to MultiAgent Systems. 2nd ed. Chichester: John Wiley & Sons, 2009. 484 p.",
    "aima": "Russell S., Norvig P. Artificial Intelligence: A Modern Approach. 4th ed. Hoboken: Pearson, 2021. " + url("https://aima.cs.berkeley.edu/", ACCESSED_NEW),
    "autogen": "Wu Q., Bansal G., Zhang J. et al. AutoGen: Enabling Next-Gen LLM Applications via Multi-Agent Conversation. arXiv:2308.08155. 2023. " + url("https://arxiv.org/abs/2308.08155", ACCESSED_OLD),
    "metagpt": "Hong S., Zhuge M., Chen J. et al. MetaGPT: Meta Programming for A Multi-Agent Collaborative Framework. arXiv:2308.00352. 2023. " + url("https://arxiv.org/abs/2308.00352", ACCESSED_OLD),
    "react": "Yao S., Zhao J., Yu D. et al. ReAct: Synergizing Reasoning and Acting in Language Models. arXiv:2210.03629. 2022. " + url("https://arxiv.org/abs/2210.03629", ACCESSED_OLD),
    "bea": "Anthropic. Building effective agents. 2024. " + url("https://www.anthropic.com/engineering/building-effective-agents", ACCESSED_OLD),
    "mars": "Anthropic. How we built our multi-agent research system. 2025. " + url("https://www.anthropic.com/engineering/multi-agent-research-system", ACCESSED_OLD),
    "lc_multi": "LangChain. Multi-agent. " + url("https://docs.langchain.com/oss/python/langchain/multi-agent", ACCESSED_OLD),
    "lg_workflows": "LangChain. Workflows and agents. " + url("https://docs.langchain.com/oss/python/langgraph/workflows-agents", ACCESSED_OLD),
    "guo_survey": "Guo T., Chen X., Wang Y. et al. Large Language Model based Multi-Agents: A Survey of Progress and Challenges. arXiv:2402.01680. 2024. " + url("https://arxiv.org/abs/2402.01680", ACCESSED_NEW),
    "mast": "Cemri M., Pan M. Z., Yang S. et al. Why Do Multi-Agent LLM Systems Fail? arXiv:2503.13657. 2025. " + url("https://arxiv.org/abs/2503.13657", ACCESSED_NEW),
    "fowler_cb": "Fowler M. CircuitBreaker. 2014. " + url("https://martinfowler.com/bliki/CircuitBreaker.html", ACCESSED_OLD),
    "whisper": "Radford A., Kim J. W., Xu T. et al. Robust Speech Recognition via Large-Scale Weak Supervision. arXiv:2212.04356. 2022. " + url("https://arxiv.org/abs/2212.04356", ACCESSED_OLD),
    "whisper_streaming": "Macháček D., Dabre R., Bojar O. Turning Whisper into Real-Time Transcription System. arXiv:2307.14743. 2023. " + url("https://arxiv.org/abs/2307.14743", ACCESSED_NEW),
    "mdn_speech": "MDN Web Docs. Web Speech API. " + url("https://developer.mozilla.org/en-US/docs/Web/API/Web_Speech_API", ACCESSED_OLD),
    "w3c_speech": "Web Speech API: Draft Community Group Report. " + url("https://webaudio.github.io/web-speech-api/", ACCESSED_OLD),
    "gcp_lang": "Google Cloud. Speech-to-Text: supported languages. " + url("https://docs.cloud.google.com/speech-to-text/docs/speech-to-text-supported-languages", ACCESSED_OLD),
    "gcp_stream": "Google Cloud. Speech-to-Text: streaming recognition. " + url("https://docs.cloud.google.com/speech-to-text/docs/streaming-recognize", ACCESSED_OLD),
    "meet_captions": "Google Meet Help. Довідка щодо субтитрів і перекладених субтитрів. " + url("https://support.google.com/meet/answer/9300310", ACCESSED_OLD),
    "meet_gemini": "Google Meet Help. Take notes for me in Google Meet. " + url("https://support.google.com/meet/answer/14754931", ACCESSED_NEW),
    "jurafsky": "Jurafsky D., Martin J. H. Speech and Language Processing. 3rd ed. draft. " + url("https://web.stanford.edu/~jurafsky/slp3/", ACCESSED_OLD),
    "nielsen": "Nielsen J. Response Times: The 3 Important Limits. Nielsen Norman Group, 1993. " + url("https://www.nngroup.com/articles/response-times-3-important-limits/", ACCESSED_NEW),
    "tail_at_scale": "Dean J., Barroso L. A. The Tail at Scale. Communications of the ACM. 2013. Vol. 56, No. 2. P. 74–80. " + url("https://research.google/pubs/the-tail-at-scale/", ACCESSED_NEW),
    "lost_middle": "Liu N. F., Lin K., Hewitt J. et al. Lost in the Middle: How Language Models Use Long Contexts. arXiv:2307.03172. 2023. " + url("https://arxiv.org/abs/2307.03172", ACCESSED_NEW),
    "ctx_eng": "Anthropic. Effective context engineering for AI agents. 2025. " + url("https://www.anthropic.com/engineering/effective-context-engineering-for-ai-agents", ACCESSED_NEW),
    "greshake": "Greshake K., Abdelnabi S., Mishra S. et al. Not what you've signed up for: Compromising Real-World LLM-Integrated Applications with Indirect Prompt Injection. arXiv:2302.12173. 2023. " + url("https://arxiv.org/abs/2302.12173", ACCESSED_NEW),
    "owasp": "OWASP Gen AI Security Project. LLM01:2025 Prompt Injection. " + url("https://genai.owasp.org/llmrisk/llm01-prompt-injection/", ACCESSED_NEW),
    "anth_pricing": "Anthropic. Pricing. " + url("https://platform.claude.com/docs/en/about-claude/pricing", ACCESSED_PRICES),
    "anth_caching": "Anthropic. Prompt caching. " + url("https://platform.claude.com/docs/en/build-with-claude/prompt-caching", ACCESSED_PRICES),
    "anth_cost": "Anthropic. Optimizing for cost and intelligence. " + url("https://platform.claude.com/docs/en/about-claude/models/optimizing-for-cost-and-intelligence", ACCESSED_PRICES),
    "texttiling": "Hearst M. A. TextTiling: Segmenting Text into Multi-paragraph Subtopic Passages. Computational Linguistics. 1997. Vol. 23, No. 1. P. 33–64. " + url("https://aclanthology.org/J97-1003/", ACCESSED_OLD),
    "galley": "Galley M., McKeown K. R., Fosler-Lussier E., Jing H. Discourse Segmentation of Multi-Party Conversation. Proceedings of the 41st Annual Meeting of the Association for Computational Linguistics. Sapporo, 2003. P. 562–569. " + url("https://aclanthology.org/P03-1071/", ACCESSED_NEW),
    "sbert": "Reimers N., Gurevych I. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. arXiv:1908.10084. 2019. " + url("https://arxiv.org/abs/1908.10084", ACCESSED_OLD),
    "multiling_sbert": "Reimers N., Gurevych I. Making Monolingual Sentence Embeddings Multilingual using Knowledge Distillation. arXiv:2004.09813. 2020. " + url("https://arxiv.org/abs/2004.09813", ACCESSED_NEW),
    "manning_ir": "Manning C. D., Raghavan P., Schütze H. Introduction to Information Retrieval. Cambridge: Cambridge University Press, 2008. " + url("https://nlp.stanford.edu/IR-book/html/htmledition/evaluation-of-unranked-retrieval-sets-1.html", ACCESSED_NEW),
    "windowdiff": "Pevzner L., Hearst M. A. A Critique and Improvement of an Evaluation Metric for Text Segmentation. Computational Linguistics. 2002. Vol. 28, No. 1. P. 19–36. " + url("https://aclanthology.org/J02-1002/", ACCESSED_NEW),
    "qmsum": "Zhong M., Yin D., Yu T. et al. QMSum: A New Benchmark for Query-based Multi-domain Meeting Summarization. arXiv:2104.05938. 2021. " + url("https://arxiv.org/abs/2104.05938", ACCESSED_NEW),
    "ami": "Carletta J. et al. The AMI Meeting Corpus: A Pre-Announcement. MLMI 2005. " + url("https://infoscience.epfl.ch/record/83204", ACCESSED_OLD),
    "meet_rest": "Google for Developers. Google Meet REST API overview. " + url("https://developers.google.com/workspace/meet/api/guides/overview", ACCESSED_OLD),
    "meet_artifacts": "Google for Developers. Meet REST API: work with artifacts. " + url("https://developers.google.com/workspace/meet/api/guides/artifacts", ACCESSED_OLD),
    "meet_media": "Google for Developers. Google Meet Media API overview. " + url("https://developers.google.com/workspace/meet/media-api/guides/overview", ACCESSED_OLD),
    "meet_addons": "Google for Developers. Google Meet add-ons SDK overview. " + url("https://developers.google.com/workspace/meet/add-ons/guides/overview", ACCESSED_OLD),
    "content_scripts": "Chrome for Developers. Content scripts. " + url("https://developer.chrome.com/docs/extensions/develop/concepts/content-scripts", ACCESSED_OLD),
    "messaging": "Chrome for Developers. Message passing. " + url("https://developer.chrome.com/docs/extensions/develop/concepts/messaging", ACCESSED_OLD),
    "sw_lifecycle": "Chrome for Developers. The extension service worker lifecycle. " + url("https://developer.chrome.com/docs/extensions/develop/concepts/service-workers/lifecycle", ACCESSED_OLD),
    "tabcapture": "Chrome for Developers. chrome.tabCapture. " + url("https://developer.chrome.com/docs/extensions/reference/api/tabCapture", ACCESSED_OLD),
    "offscreen": "Chrome for Developers. chrome.offscreen. " + url("https://developer.chrome.com/docs/extensions/reference/api/offscreen", ACCESSED_OLD),
    "permissions": "Chrome for Developers. Declare permissions. " + url("https://developer.chrome.com/docs/extensions/develop/concepts/declare-permissions", ACCESSED_OLD),
    "chrome_storage": "Chrome for Developers. chrome.storage. " + url("https://developer.chrome.com/docs/extensions/reference/api/storage", ACCESSED_NEW),
    "mdn_mutation": "MDN Web Docs. MutationObserver. " + url("https://developer.mozilla.org/en-US/docs/Web/API/MutationObserver", ACCESSED_NEW),
    "mdn_shadow": "MDN Web Docs. Using shadow DOM. " + url("https://developer.mozilla.org/en-US/docs/Web/API/Web_components/Using_shadow_DOM", ACCESSED_NEW),
    "rfc6455": "Fette I., Melnikov A. The WebSocket Protocol: RFC 6455. IETF, 2011. " + url("https://www.rfc-editor.org/rfc/rfc6455", ACCESSED_OLD),
    "pydantic": "Pydantic. Documentation. " + url("https://docs.pydantic.dev/latest/", ACCESSED_NEW),
    "asyncio": "Python Software Foundation. asyncio — Asynchronous I/O. " + url("https://docs.python.org/3/library/asyncio.html", ACCESSED_PRICES),
    "fastapi_ws": "FastAPI. WebSockets. " + url("https://fastapi.tiangolo.com/advanced/websockets/", ACCESSED_NEW),
    "anth_messages": "Anthropic. Messages API. " + url("https://platform.claude.com/docs/en/api/messages", ACCESSED_NEW),
    "anth_structured": "Anthropic. Structured outputs. " + url("https://platform.claude.com/docs/en/docs/build-with-claude/structured-outputs", ACCESSED_OLD),
    "vite": "Vite. Building for Production. " + url("https://vite.dev/guide/build", ACCESSED_NEW),
    "pytest": "pytest: helps you write better programs. " + url("https://docs.pytest.org/en/stable/", ACCESSED_NEW),
}

CITE = re.compile(r"\[@([a-z0-9_]+)\]")


def resolve(text: str) -> tuple[str, list[str]]:
    order: list[str] = []

    def sub(m: re.Match[str]) -> str:
        key = m.group(1)
        if key not in REFS:
            raise KeyError(f"невідоме джерело: {key}")
        if key not in order:
            order.append(key)
        return f"[{order.index(key) + 1}]"

    return CITE.sub(sub, text), order


def density_report(text: str) -> list[str]:
    """Позиції, де у вікні PAGE_CHARS символів основного тексту більше дозволеного різних джерел."""
    plain = "\n".join(l for l in text.split("\n") if not l.startswith("!["))
    cites = [(m.start(), m.group(0)) for m in re.finditer(r"\[(\d+)\]", plain)]
    problems, last = [], -PAGE_CHARS
    for i, (pos, _) in enumerate(cites):
        window = {c for p, c in cites[i:] if p - pos <= PAGE_CHARS}
        if len(window) > MAX_SOURCES_PER_PAGE and pos - last > PAGE_CHARS // 2:
            problems.append(f"{len(window)} джерел біля: «{plain[pos - 60:pos + 20].strip()}»")
            last = pos
    return problems


def heading(doc, text: str, level: int, first: bool) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1 and not first:
        p.paragraph_format.page_break_before = True
    bd.add_inline(p, text, bold=True)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
    structural = level == 1 and not text[0].isdigit()
    bd.fmt(p, WD_ALIGN_PARAGRAPH.CENTER if structural else WD_ALIGN_PARAGRAPH.LEFT, indent=not structural,
           before=0 if level == 1 else 12, after=12 if level == 1 else 6, keep=True)
    if level == 1:
        p.paragraph_format.page_break_before = not first


def toc(doc) -> None:
    p = doc.add_paragraph()
    r = p.add_run("ЗМІСТ"); bd.set_font(r, 14, True)
    bd.fmt(p, WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=12)
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, txt in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        run._r.append(el)
    r2 = p.add_run("Оновіть поле змісту: клацніть тут і натисніть F9."); bd.set_font(r2, 12, italic=True)
    r3 = p.add_run(); el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), "end"); r3._r.append(el)
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)


def picture(doc, path: Path, caption: str) -> None:
    with Image.open(path) as im:
        w, h = im.size
    width = min(MAX_IMAGE_WIDTH_CM, MAX_IMAGE_HEIGHT_CM * w / h)
    p = doc.add_paragraph(); bd.fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 0, True)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Cm(width))
    cp = doc.add_paragraph(); bd.add_inline(cp, caption)
    bd.fmt(cp, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 12)


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bd.borders(t)
    for i, row in enumerate([header] + body):
        for j, txt in enumerate(row[:len(header)]):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            bd.add_inline(p, txt.strip(), size=TABLE_FONT_PT, bold=(i == 0))
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            if i == 0:
                bd.shade(c, "E7E6E6")
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = 0.7


def code_block(doc, lines: list[str]) -> None:
    for i, ln in enumerate(lines):
        p = doc.add_paragraph()
        r = p.add_run(ln or " "); bd.set_font(r, 10.5, mono=True)
        pf = p.paragraph_format
        pf.line_spacing = 1.0; pf.left_indent = Cm(1.25)
        pf.space_before = Pt(6 if i == 0 else 0); pf.space_after = Pt(6 if i == len(lines) - 1 else 0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT


def build_docx(text: str, out: Path) -> None:
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    # ДСТУ 3008:2015: лівий, верхній і нижній береги не менше 20 мм, правий — не менше 10 мм.
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2.5), Cm(1.5), Cm(2.0), Cm(2.0)
    st = doc.styles["Normal"]; st.font.name = bd.FONT; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), bd.FONT)
    for level in (1, 2, 3):
        hs = doc.styles[f"Heading {level}"]
        hs.font.name = bd.FONT; hs.font.size = Pt(14); hs.font.bold = True
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), bd.FONT)
        hs.element.rPr.rFonts.set(qn("w:ascii"), bd.FONT); hs.element.rPr.rFonts.set(qn("w:hAnsi"), bd.FONT)
    bd.page_number_footer(doc)
    toc(doc)

    lines = text.split("\n")
    i, first = 0, False
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("# "):
            heading(doc, ln[2:], 1, first); first = False
        elif ln.startswith("## Висновки до розділу"):
            # не заголовок: у зміст потрапляють лише пункти плану роботи
            p = doc.add_paragraph(); bd.add_inline(p, ln[3:], bold=True)
            bd.fmt(p, WD_ALIGN_PARAGRAPH.LEFT, True, 12, 6, True)
        elif ln.startswith("## "):
            heading(doc, ln[3:], 2, False)
        elif ln.startswith("### "):
            heading(doc, ln[4:], 3, False)
        elif ln.startswith("```"):
            block = []; i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i]); i += 1
            code_block(doc, block)
        elif ln.startswith("$$ "):
            p = doc.add_paragraph(); bd.add_inline(p, ln[3:].strip())
            bd.fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 6)
        elif ln.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.+)\)", ln)
            picture(doc, (REPORT / m.group(2)).resolve(), m.group(1))
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = lines[i].strip().strip("|").split("|")
                if not all(re.fullmatch(r"\s*-+\s*", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows); continue
        elif re.match(r"^(\d+)\.\s", ln) or ln.startswith("- "):
            m = re.match(r"^(\d+\.)\s(.*)", ln)
            p = doc.add_paragraph()
            bd.add_inline(p, f"{m.group(1)} {m.group(2)}" if m else "– " + ln[2:])
            bd.fmt(p, indent=True)
        else:
            p = doc.add_paragraph(); bd.add_inline(p, ln)
            cap = ln.startswith("**Таблиця")
            bd.fmt(p, WD_ALIGN_PARAGRAPH.LEFT if cap else WD_ALIGN_PARAGRAPH.JUSTIFY, not cap, 6 if cap else 0, 3 if cap else 0, cap)
        i += 1
    doc.save(out)


def main() -> None:
    global REPORT
    folder, out_name = VARIANTS[sys.argv[1] if len(sys.argv) > 1 else "practice"]
    REPORT = HERE.parent / folder
    parts = PRACTICE_PARTS if folder == "report" else PARTS
    raw = "\n\n".join((REPORT / name).read_text(encoding="utf-8").strip() for name in parts)
    text, order = resolve(raw)
    unused = [k for k in REFS if k not in order]
    refs = "\n".join(f"{n}. {REFS[k]}" for n, k in enumerate(order, start=1))
    text += "\n\n# ПЕРЕЛІК ВИКОРИСТАНИХ ДЖЕРЕЛ\n\n" + refs + "\n"
    (REPORT / f"{out_name}.md").write_text(text, encoding="utf-8")
    build_docx(text, REPORT / f"{out_name}.docx")
    words = len(re.findall(r"\w+", raw))
    print(f"джерел: {len(order)}; не використано: {unused}; слів: {words}")
    for problem in density_report(text.split("# ПЕРЕЛІК ВИКОРИСТАНИХ ДЖЕРЕЛ")[0]):
        print("ЩІЛЬНІСТЬ:", problem)


if __name__ == "__main__":
    main()
